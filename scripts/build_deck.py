"""Builds Laya_Project_Overview.pptx: a technical briefing deck covering
what Laya is, how it compares to Jev, the heuristic/laya-mlx/Jev three-way
split, the two demo games, measured results, and /decide API examples.

Usage: uv run python scripts/build_deck.py
"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

# ---------------------------------------------------------------- palette --

NAVY = RGBColor(0x12, 0x19, 0x2B)
OFFWHITE = RGBColor(0xF6, 0xF5, 0xF1)
CARD = RGBColor(0xFF, 0xFF, 0xFF)
TEAL = RGBColor(0x16, 0xA3, 0x94)
GREEN = RGBColor(0x2F, 0x8F, 0x5B)
AMBER = RGBColor(0xC9, 0x7A, 0x2E)
SOFTDARK = RGBColor(0x30, 0x39, 0x4A)
MUTED = RGBColor(0x6A, 0x71, 0x79)
LIGHT_TEXT = RGBColor(0xF6, 0xF5, 0xF1)
MUTED_LIGHT = RGBColor(0xBF, 0xD8, 0xD5)
BORDER = RGBColor(0xE3, 0xE6, 0xE4)

HEAD_FONT = "Arial"
BODY_FONT = "Arial"
MONO_FONT = "Consolas"

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)


def new_deck():
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    return prs


def blank_slide(prs, bg=OFFWHITE):
    layout = prs.slide_layouts[6]  # blank
    slide = prs.slides.add_slide(layout)
    bg_shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg_shape.fill.solid()
    bg_shape.fill.fore_color.rgb = bg
    bg_shape.line.fill.background()
    bg_shape.shadow.inherit = False
    # send to back
    sp = bg_shape._element
    sp.getparent().remove(sp)
    slide.shapes._spTree.insert(2, sp)
    return slide


def add_text(slide, x, y, w, h, text, size, color, bold=False, italic=False,
             font=BODY_FONT, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
             line_spacing=1.0, letter_tracking=None):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    lines = text.split("\n")
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = line_spacing
        run = p.add_run()
        run.text = line
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.italic = italic
        run.font.name = font
        run.font.color.rgb = color
    return box


def add_bullets(slide, x, y, w, h, items, size, color, font=BODY_FONT,
                 line_spacing=1.15, space_after=10, bullet_char="—", bold_first=False):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.line_spacing = line_spacing
        p.space_after = Pt(space_after)
        run = p.add_run()
        run.text = f"{bullet_char}  {item}"
        run.font.size = Pt(size)
        run.font.name = font
        run.font.color.rgb = color
    return box


def add_rect(slide, x, y, w, h, fill=None, line_color=None, line_width=None, radius=None):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(shape_type, x, y, w, h)
    if radius:
        try:
            shape.adjustments[0] = radius
        except Exception:
            pass
    if fill is None:
        shape.fill.background()
    else:
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill
    if line_color is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line_color
        shape.line.width = line_width or Pt(1)
    shape.shadow.inherit = False
    return shape


def set_cell(cell, text, size, color, bold=False, bg=None, font=BODY_FONT,
             align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE):
    cell.vertical_anchor = anchor
    cell.margin_left = Pt(10)
    cell.margin_right = Pt(10)
    cell.margin_top = Pt(6)
    cell.margin_bottom = Pt(6)
    if bg is not None:
        cell.fill.solid()
        cell.fill.fore_color.rgb = bg
    else:
        cell.fill.solid()
        cell.fill.fore_color.rgb = CARD
    tf = cell.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.name = font
    run.font.color.rgb = color


def add_table(slide, x, y, w, h, headers, rows, col_widths, header_bg=NAVY,
              header_color=LIGHT_TEXT, body_color=SOFTDARK, size=16,
              header_size=15, row_height=None, highlight_col=None,
              highlight_color=None):
    n_rows = len(rows) + 1
    n_cols = len(headers)
    gshape = slide.shapes.add_table(n_rows, n_cols, x, y, w, h)
    table = gshape.table
    for i, cw in enumerate(col_widths):
        table.columns[i].width = cw
    if row_height:
        for r in table.rows:
            r.height = row_height
    for j, htext in enumerate(headers):
        set_cell(table.cell(0, j), htext, header_size, header_color, bold=True,
                 bg=header_bg, align=PP_ALIGN.LEFT if j == 0 else PP_ALIGN.CENTER)
    for i, row in enumerate(rows, start=1):
        for j, val in enumerate(row):
            bg = CARD if i % 2 == 1 else RGBColor(0xF1, 0xF0, 0xEB)
            color = body_color
            bold = False
            if highlight_col is not None and j == highlight_col:
                color = highlight_color or body_color
                bold = True
            set_cell(table.cell(i, j), val, size, color, bold=bold, bg=bg,
                     align=PP_ALIGN.LEFT if j == 0 else PP_ALIGN.CENTER)
    return table


def footer(slide, label, page_no):
    add_text(slide, Inches(0.7), Inches(7.08), Inches(6), Inches(0.35),
              label, 11, MUTED, font=BODY_FONT)
    add_text(slide, Inches(12.3), Inches(7.08), Inches(0.6), Inches(0.35),
              str(page_no), 11, MUTED, align=PP_ALIGN.RIGHT)


def eyebrow(slide, text, color=TEAL):
    add_text(slide, Inches(0.7), Inches(0.55), Inches(8), Inches(0.4),
              text.upper(), 14, color, bold=True, font=BODY_FONT)


def title(slide, text, color=NAVY, y=0.95, size=38):
    add_text(slide, Inches(0.7), Inches(y), Inches(11.9), Inches(0.9),
              text, size, color, bold=True, font=HEAD_FONT)


def card(slide, x, y, w, h, heading, lines, accent=TEAL, heading_color=NAVY,
          body_color=SOFTDARK, heading_size=19, body_size=14):
    add_rect(slide, x, y, w, h, fill=CARD, line_color=BORDER, line_width=Pt(1), radius=0.06)
    add_rect(slide, x, y, Inches(0.07), h, fill=accent)
    add_text(slide, x + Inches(0.3), y + Inches(0.22), w - Inches(0.55), Inches(0.5),
              heading, heading_size, heading_color, bold=True, font=HEAD_FONT)
    add_bullets(slide, x + Inches(0.3), y + Inches(0.75), w - Inches(0.55), h - Inches(0.95),
                lines, body_size, body_color, line_spacing=1.2, space_after=6, bullet_char="·")


# ------------------------------------------------------------------ build --

prs = new_deck()

# 1. Cover ------------------------------------------------------------------
s = blank_slide(prs, bg=NAVY)
add_rect(s, 0, Inches(6.55), SLIDE_W, Inches(0.05), fill=TEAL)
add_text(s, Inches(0.9), Inches(2.55), Inches(6), Inches(0.5),
          "TECHNICAL BRIEFING", 15, TEAL, bold=True, font=BODY_FONT)
add_text(s, Inches(0.9), Inches(3.0), Inches(11.2), Inches(1.8),
          "Laya: Typed Decisions\nWithout Text Generation", 46, LIGHT_TEXT, bold=True,
          font=HEAD_FONT, line_spacing=1.05)
add_text(s, Inches(0.9), Inches(4.75), Inches(10.5), Inches(0.8),
          "An open-source System-1 decision engine — our own heuristic engine,\n"
          "laya-mlx, and Jev, compared on real measured results.",
          19, MUTED_LIGHT, font=BODY_FONT, line_spacing=1.3)
add_text(s, Inches(0.9), Inches(6.75), Inches(8), Inches(0.4),
          "Engineering & Product  ·  Laya Demo Project", 13, MUTED_LIGHT)

# 2. What is Laya -------------------------------------------------------------
s = blank_slide(prs)
eyebrow(s, "Background")
title(s, "What Is Laya?")
add_text(s, Inches(0.7), Inches(1.85), Inches(11.5), Inches(0.9),
          "An open-source, Apache 2.0 model that answers typed questions directly —\n"
          "bool / enum / number — in a single forward pass. No text generation.",
          20, SOFTDARK, line_spacing=1.3)
add_bullets(s, Inches(0.7), Inches(2.9), Inches(11.7), Inches(2.6), [
    "Open alternative to Jev (TypeSafe AI's commercial “System One” API)",
    "Three checkpoints, 322M–421M parameters — English + multilingual",
    "Ships as laya-mlx for native Apple Silicon inference (Metal/MLX)",
    "Target workloads: invoice processing, ticket routing, moderation, guardrails",
], 16, SOFTDARK, space_after=12, line_spacing=1.1)

chip_labels = ["Apache 2.0", "3 checkpoints", "Runs on Apple Silicon", "$0 per call, self-hosted"]
cx = Inches(0.7)
cy = Inches(6.1)
for lbl in chip_labels:
    w = Inches(0.35 + 0.1 * len(lbl))
    add_rect(s, cx, cy, w, Inches(0.55), fill=RGBColor(0xEA, 0xF6, 0xF4), radius=0.5)
    add_text(s, cx, cy, w, Inches(0.55), lbl, 14, GREEN, bold=True, align=PP_ALIGN.CENTER,
              anchor=MSO_ANCHOR.MIDDLE)
    cx = Emu(cx + w + Inches(0.25))
footer(s, "Laya Project Overview", 2)

# 3. Jev vs Laya --------------------------------------------------------------
s = blank_slide(prs)
eyebrow(s, "Positioning")
title(s, "Jev vs. Laya")
headers = ["", "Jev (commercial)", "Laya (open-source)"]
rows = [
    ["License", "Closed, pay-per-token", "Apache 2.0, free"],
    ["Data privacy", "Requests leave your infra", "Fully local — nothing leaves your machine"],
    ["Latency", "Network round-trip to vendor", "In-process, no network hop"],
    ["Control", "Black-box decisions", "Fully auditable code"],
    ["Vendor lock-in", "Tied to vendor pricing & roadmap", "None — you own the code"],
]
add_table(s, Inches(0.7), Inches(1.9), Inches(11.9), Inches(4.3), headers, rows,
          col_widths=[Inches(2.3), Inches(4.8), Inches(4.8)], size=16, header_size=16)
add_text(s, Inches(0.7), Inches(6.45), Inches(11.5), Inches(0.5),
          "Jev likely has an edge in raw accuracy from a managed, trained model — the rest of this "
          "deck measures how far an open alternative gets.", 14, MUTED, italic=True)
footer(s, "Laya Project Overview", 3)

# 4. Three approaches ----------------------------------------------------------
s = blank_slide(prs)
eyebrow(s, "This Project")
title(s, "Three Ways to Make a Typed Decision")
add_text(s, Inches(0.7), Inches(1.75), Inches(11), Inches(0.5),
          "What we actually built and tested — not a hypothetical comparison.",
          17, MUTED)
headers = ["", "Heuristic (ours)", "laya-mlx", "Jev"]
rows = [
    ["Mechanism", "Hand-written word-overlap scoring", "Trained decision transformer (322–421M)", "Trained model (proprietary)"],
    ["Training data", "None — no model at all", "General-purpose typed-decision tasks", "Unknown, vendor-controlled"],
    ["Setup", "Zero — nothing to install", "One-time model download (~several 100MB)", "API key + network access"],
    ["Used in this project", "Yes — powers both demo games today", "Yes — benchmarked side by side", "No — reference only, not integrated"],
]
add_table(s, Inches(0.7), Inches(2.4), Inches(11.9), Inches(3.9), headers, rows,
          col_widths=[Inches(2.1), Inches(3.5), Inches(3.5), Inches(2.8)], size=14.5, header_size=15)
footer(s, "Laya Project Overview", 4)

# 5. Pattern --------------------------------------------------------------------
s = blank_slide(prs)
eyebrow(s, "How The Demos Work")
title(s, "Same Pattern, Every Flow")
steps = [
    ("Game State", "Real math:\npositions, health,\nammo, offsets"),
    ("English Sentence", "Templated text,\ne.g. “enemy\nspotted left”"),
    ("/decide", "Typed enum\nquestion sent to\nheuristic or laya-mlx"),
    ("Action", "Mapped to a real\ngame input\n(turn, shoot, move)"),
]
box_w = Inches(2.6)
gap = Inches(0.55)
total_w = box_w * 4 + gap * 3
start_x = Emu(int((SLIDE_W - total_w) / 2))
y = Inches(2.6)
box_h = Inches(2.3)
x = start_x
for i, (head, body) in enumerate(steps):
    accent = TEAL if i % 2 == 0 else GREEN
    add_rect(s, x, y, box_w, box_h, fill=CARD, line_color=BORDER, line_width=Pt(1), radius=0.08)
    add_rect(s, x, y, box_w, Inches(0.08), fill=accent)
    add_text(s, x + Inches(0.2), y + Inches(0.3), box_w - Inches(0.4), Inches(0.6),
              head, 18, NAVY, bold=True, font=HEAD_FONT)
    add_text(s, x + Inches(0.2), y + Inches(0.95), box_w - Inches(0.4), Inches(1.2),
              body, 13.5, SOFTDARK, line_spacing=1.25)
    if i < 3:
        arrow_x = Emu(x + box_w)
        arrow = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, arrow_x, Emu(int(y + box_h / 2 - Inches(0.12))),
                                     gap, Inches(0.24))
        arrow.fill.solid()
        arrow.fill.fore_color.rgb = MUTED
        arrow.line.fill.background()
        arrow.shadow.inherit = False
    x = Emu(x + box_w + gap)
add_text(s, Inches(0.7), Inches(5.6), Inches(11.5), Inches(0.6),
          "Neither engine ever sees pixels, coordinates, or angles — only a short sentence.",
          18, SOFTDARK, italic=True, align=PP_ALIGN.CENTER)
footer(s, "Laya Project Overview", 5)

# 6. Two games --------------------------------------------------------------------
s = blank_slide(prs)
eyebrow(s, "Demo Games")
title(s, "Two Demo Games, Same Decision Loop")
card(s, Inches(0.7), Inches(2.0), Inches(5.75), Inches(4.5), "play_doom.py — Real Doom",
     [
         "VizDoom scenario: defend_the_center",
         "Enemies close in from all sides",
         "3 options: turn_left / turn_right / shoot",
         "Decision made ~20× per second",
         "Fully autonomous — no player input",
     ], accent=TEAL, body_size=15)
card(s, Inches(6.85), Inches(2.0), Inches(5.75), Inches(4.5), "shooting_range.py — Shooting Range",
     [
         "Custom pygame turret + monster",
         "Player moves the monster (arrow keys)",
         "5 options: turn_left / right / up / down / shoot",
         "Decision made every 2 frames",
         "Isolated view of targeting logic alone",
     ], accent=GREEN, body_size=15)
footer(s, "Laya Project Overview", 6)

# 7. Results ------------------------------------------------------------------------
s = blank_slide(prs)
eyebrow(s, "Measured Results")
title(s, "Heuristic vs. laya-mlx")
add_text(s, Inches(0.7), Inches(1.75), Inches(11.5), Inches(0.45),
          "All numbers from real runs against this codebase — not estimates.", 15, MUTED)
headers = ["Flow", "Metric", "Heuristic", "laya-mlx"]
rows = [
    ["/decide API", "Accuracy (6 cases)", "100% (6/6)", "67% (4/6)"],
    ["/decide API", "Latency (steady state)", "~1.5 ms", "~30 ms"],
    ["Doom", "Shoot accuracy*", "100%", "~22%"],
    ["Doom", "Avg. survival", "~1,106 ticks", "~375 ticks"],
    ["Shooting Range", "Decision accuracy", "100% (121/121)", "33% (12/36)"],
]
add_table(s, Inches(0.7), Inches(2.35), Inches(11.9), Inches(3.7), headers, rows,
          col_widths=[Inches(2.9), Inches(3.1), Inches(3.0), Inches(2.9)],
          size=16, header_size=15.5, highlight_col=2, highlight_color=GREEN)
add_text(s, Inches(0.7), Inches(6.25), Inches(11.5), Inches(0.5),
          "*Shoot accuracy = fraction of “shoot” decisions where a live monster was actually "
          "centered. Checkpoint used: aac6fef/laya-typed-decisions-mlx (best of 3 tested).",
          12.5, MUTED, italic=True)
footer(s, "Laya Project Overview", 7)

# 8. API examples ---------------------------------------------------------------------
s = blank_slide(prs)
eyebrow(s, "The /decide API")
title(s, "Three Question Types, One Endpoint")

examples = [
    ("bool", '{\n  "question": "Confirm this account\n   is safe and valid",\n  "type": "bool"\n}',
     '{\n  "answer": true,\n  "confidence": 0.95,\n  "latency_ms": 0.04\n}'),
    ("enum", '{\n  "question": "Customer wants a\n   refund for late delivery",\n  "type": "enum",\n  "options": [...]\n}',
     '{\n  "answer": "refund",\n  "confidence": 0.70,\n  "latency_ms": 0.03\n}'),
    ("number", '{\n  "question": "Rate this 1 to 5,\n   I’d say a 4",\n  "type": "number",\n  "min_value": 1, "max_value": 5\n}',
     '{\n  "answer": 4.0,\n  "confidence": 0.90,\n  "latency_ms": 0.01\n}'),
]
col_w = Inches(3.9)
gap = Inches(0.2)
x = Inches(0.7)
y = Inches(1.9)
for name, req, resp in examples:
    add_rect(s, x, y, col_w, Inches(0.5), fill=TEAL, radius=0.12)
    add_text(s, x, y, col_w, Inches(0.5), name, 18, LIGHT_TEXT, bold=True,
              font=MONO_FONT, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    add_rect(s, x, Emu(y + Inches(0.6)), col_w, Inches(2.1), fill=NAVY, radius=0.06)
    add_text(s, Emu(x + Inches(0.18)), Emu(y + Inches(0.78)), col_w - Inches(0.36), Inches(1.8),
              req, 11, MUTED_LIGHT, font=MONO_FONT, line_spacing=1.2)
    add_text(s, x, Emu(y + Inches(2.85)), col_w, Inches(0.3), "↓ response", 12, MUTED,
              align=PP_ALIGN.CENTER, italic=True)
    add_rect(s, x, Emu(y + Inches(3.2)), col_w, Inches(1.5), fill=RGBColor(0xEA, 0xF6, 0xF4), radius=0.06)
    add_text(s, Emu(x + Inches(0.18)), Emu(y + Inches(3.38)), col_w - Inches(0.36), Inches(1.2),
              resp, 12, SOFTDARK, font=MONO_FONT, line_spacing=1.25)
    x = Emu(x + col_w + gap)
footer(s, "Laya Project Overview", 8)

# 9. Live demo ------------------------------------------------------------------------
s = blank_slide(prs, bg=NAVY)
add_text(s, Inches(0.9), Inches(2.5), Inches(8), Inches(0.5),
          "LIVE DEMO", 16, TEAL, bold=True)
add_text(s, Inches(0.9), Inches(3.0), Inches(10), Inches(1.0),
          "Seeing it run", 46, LIGHT_TEXT, bold=True, font=HEAD_FONT)
add_bullets(s, Inches(0.9), Inches(4.2), Inches(9.5), Inches(2.0), [
    "Doom gameplay on the heuristic backend",
    "Shooting range on the heuristic backend",
    "Swagger UI: /decide calls on both backends, live",
], 20, MUTED_LIGHT, space_after=16, bullet_char="→")
footer_text_box = add_text(s, Inches(0.9), Inches(7.08), Inches(8), Inches(0.35),
          "Laya Project Overview", 11, MUTED_LIGHT)

# 10. Bottom line --------------------------------------------------------------------
s = blank_slide(prs)
eyebrow(s, "Recommendation")
title(s, "Bottom Line")
card(s, Inches(0.7), Inches(1.9), Inches(11.9), Inches(1.5), "Today",
     ["The heuristic engine wins on both speed and accuracy for these demos — "
      "because it was built specifically for this exact wording."],
     accent=GREEN, body_size=16, heading_size=18)
card(s, Inches(0.7), Inches(3.55), Inches(11.9), Inches(1.5), "laya-mlx",
     ["Proves the real typed-decision contract and genuine model inference — but it's "
      "out-of-domain on Doom-style spatial text, not a drop-in upgrade yet."],
     accent=AMBER, body_size=16, heading_size=18)
card(s, Inches(0.7), Inches(5.2), Inches(11.9), Inches(1.5), "Next step",
     ["Keep the heuristic for demo reliability. Evaluate laya-mlx on-domain "
      "(tickets, approvals, moderation) before any production decision — "
      "dedicated GPU hosting helps throughput at scale, not accuracy."],
     accent=TEAL, body_size=16, heading_size=18)
footer(s, "Laya Project Overview", 10)

# 11. Closing ---------------------------------------------------------------------------
s = blank_slide(prs, bg=NAVY)
add_text(s, Inches(0.9), Inches(3.1), Inches(10), Inches(1.0),
          "Questions?", 48, LIGHT_TEXT, bold=True, font=HEAD_FONT)
add_text(s, Inches(0.9), Inches(4.1), Inches(9), Inches(0.6),
          "Repo: laya-demo  ·  Report: LAYA_MLX_VS_HEURISTIC.md", 16, MUTED_LIGHT)

prs.save("/Users/vighnesh/Practice/laya-demo/Laya_Project_Overview.pptx")
print("Saved Laya_Project_Overview.pptx")
